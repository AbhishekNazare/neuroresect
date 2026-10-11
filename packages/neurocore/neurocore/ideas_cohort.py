"""Explicit source-table mapping for retrospective IDEAS Lausanne-36 experiments."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from neurocore.ideas import index_network_archive, read_network_matrix, sha256
from neurocore.io import write_json

CLINICAL_COLUMNS = ("Sex", "Binned_Onset_Age", "Binned_Age_at_Scan", "Number_ASMs")
ENDPOINT = {
    "source_column": "ILAE_Year1",
    "followup_months": 12,
    "positive_class": "ILAE class 1",
    "negative_class": "ILAE classes 2–6",
    "definition": "Reported first-year ILAE class 1 versus classes 2–6; NA excluded",
    "reference": "https://www.ilae.org/files/ilaeGuideline/New-Classification-of-OutcomeFollowing-Epilepsy-Surgery-2001.pdf",
}


def subject_id(value: str) -> str:
    if not value.isdecimal() or int(value) <= 0:
        raise ValueError(f"Invalid numeric source subject ID: {value!r}")
    return f"sub-{int(value)}"


def read_subject_table(path: str | Path) -> tuple[dict, list[str]]:
    records: dict[str, dict] = {}
    duplicates: list[str] = []
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
            raise ValueError("Clinical table requires unique column names")
        if "ID" not in reader.fieldnames:
            raise ValueError("Clinical table requires ID")
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ValueError("Ragged clinical table")
            identifier = subject_id(row["ID"].strip())
            if identifier in records:
                if row != records[identifier]:
                    raise ValueError(f"Conflicting clinical rows for {identifier}")
                duplicates.append(identifier)
            records[identifier] = row
    return records, duplicates


def canonical_label(value: str) -> str:
    value = value.strip().strip("'")
    if value.startswith("ctx-lh-"):
        return "l." + value.removeprefix("ctx-lh-")
    if value.startswith("ctx-rh-"):
        return "r." + value.removeprefix("ctx-rh-")
    return value


def read_labels(path: str | Path) -> list[str]:
    with Path(path).open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            sheet = book["scale36"]
            labels = []
            for row in sheet.iter_rows():
                if row[0].value is None:
                    continue
                if row[0].data_type == "f" or not isinstance(row[0].value, str):
                    raise ValueError("Atlas labels must be literal strings")
                labels.append(canonical_label(row[0].value))
        finally:
            book.close()
    if len(labels) != 82 or len(set(labels)) != 82:
        raise ValueError("Lausanne-36 requires exactly 82 unique region labels")
    if any(not name.startswith(("Left-", "Right-", "l.", "r.")) for name in labels):
        raise ValueError("Unknown hemisphere naming convention")
    return labels


def read_resections(path: str | Path, labels: list[str], unit: str) -> dict:
    if unit not in ("fraction", "percent"):
        raise ValueError("Explicit resection unit must be fraction or percent")
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        table = list(csv.reader(stream))
    if not table or len(table[0]) < 2:
        raise ValueError("Resection table is empty")
    identifiers = [subject_id(value) for value in table[0][1:]]
    if len(set(identifiers)) != len(identifiers):
        raise ValueError("Duplicate resection subject columns")
    source_labels = [canonical_label(row[0]) for row in table[1:] if row]
    if len(source_labels) != len(set(source_labels)) or set(source_labels) != set(labels):
        raise ValueError("Resection region labels do not match the 82-region atlas")
    if any(len(row) != len(identifiers) + 1 for row in table[1:]):
        raise ValueError("Ragged resection table")
    values = np.asarray([row[1:] for row in table[1:]], dtype=float)
    values = values / (100 if unit == "percent" else 1)
    if np.isinf(values).any() or (values < 0).any() or (values > 1).any():
        raise ValueError("Resection fractions must be finite values within [0,1]")
    ordered = values[[source_labels.index(label) for label in labels], :]
    return {
        identifier: ordered[:, index].tolist() if np.isfinite(ordered[:, index]).all() else None
        for index, identifier in enumerate(identifiers)
    }


def map_outcome(value: str) -> int | None:
    if value in ("NA", ""):
        return None
    if value not in ("1", "2", "3", "4", "5", "6"):
        raise ValueError(f"Unknown ILAE outcome value: {value!r}")
    return int(value == "1")


def build_ideas_cohort(
    patients_path: str | Path,
    controls_path: str | Path,
    labels_path: str | Path,
    resections_path: str | Path,
    archive_path: str | Path,
    output: str | Path,
    resection_unit: str,
) -> dict:
    patients, duplicate_patients = read_subject_table(patients_path)
    controls, duplicate_controls = read_subject_table(controls_path)
    if patients.keys() & controls.keys():
        raise ValueError("Patient/control identifiers overlap")
    for row in patients.values():
        if not all(column in row for column in (*CLINICAL_COLUMNS, "ILAE_Year1")):
            raise ValueError("Required clinical columns are missing")
        map_outcome(row["ILAE_Year1"])
    labels = read_labels(labels_path)
    resections = read_resections(resections_path, labels, resection_unit)
    inventory = index_network_archive(archive_path)
    selected = [
        row
        for row in inventory["matrices"]
        if row["atlas"] == "Lausanne-36" and row["measure"] == "Count"
    ]
    if not selected:
        raise ValueError("No Lausanne-36 Count matrices")
    seen, rows = set(), []
    for matrix_row in selected:
        identifier = matrix_row["subject_id"]
        if identifier in seen:
            raise ValueError("Multiple sessions require an explicit selection policy")
        seen.add(identifier)
        clinical = patients.get(identifier)
        group = (
            "surgical_patient"
            if clinical
            else "healthy_control"
            if identifier in controls
            else "unknown"
        )
        target = map_outcome(clinical["ILAE_Year1"]) if clinical else None
        fractions = resections.get(identifier)
        reasons = []
        if group != "surgical_patient":
            reasons.append("not_surgical_patient")
        if target is None:
            reasons.append("missing_year1_outcome")
        if fractions is None:
            reasons.append(
                "incomplete_resection" if identifier in resections else "missing_resection"
            )
        elif not any(fractions):
            reasons.append("zero_resection")
        matrix = read_network_matrix(archive_path, matrix_row["member"])
        if matrix["region_count"] != len(labels):
            raise ValueError(f"Matrix/label dimension mismatch for {identifier}")
        rows.append(
            {
                "subject_id": identifier,
                "session_id": matrix_row["session_id"],
                "group": group,
                "target": target,
                "source_outcome": clinical["ILAE_Year1"] if clinical else None,
                "clinical": {key: clinical[key] for key in CLINICAL_COLUMNS} if clinical else {},
                "resection_fractions": fractions,
                "matrix_member": matrix_row["member"],
                "matrix_csv_sha256": matrix["source_csv_sha256"],
                "eligible": not reasons,
                "exclusion_reasons": reasons,
            }
        )
    eligible = [row for row in rows if row["eligible"]]
    report = {
        "schema_version": 1,
        "dataset_id": "ideas-ii-lausanne36-count-year1-v1",
        "endpoint": ENDPOINT,
        "atlas": "Lausanne-36",
        "measure": "Count",
        "tractography": selected[0]["tractography"],
        "resection_input_unit": resection_unit,
        "region_labels": labels,
        "region_name_mapping_verified": True,
        "anatomical_coordinates_available": False,
        "archive_path": str(Path(archive_path).resolve()),
        "source_hashes": {
            name: sha256(Path(path))
            for name, path in {
                "patients": patients_path,
                "controls": controls_path,
                "labels": labels_path,
                "resections": resections_path,
                "network_archive": archive_path,
            }.items()
        },
        "exact_duplicate_rows_collapsed": {
            "patients": duplicate_patients,
            "controls": duplicate_controls,
        },
        "source_patient_count": len(patients),
        "source_control_count": len(controls),
        "network_subject_count": len(rows),
        "network_group_counts": dict(Counter(row["group"] for row in rows)),
        "eligible_count": len(eligible),
        "class_counts": dict(Counter(str(row["target"]) for row in eligible)),
        "exclusion_counts": dict(
            Counter(reason for row in rows for reason in row["exclusion_reasons"])
        ),
        "subjects": rows,
        "limitations": [
            "Retrospective use of actual resections; not prospective validation.",
            "Region names match after explicit ctx-lh/ctx-rh normalization; patient imaging registration is not verified.",
            "Resection unit is explicitly configured; source website says percentage but supplied values range 0–1.",
            "Age bins are retained as categories; exact ages and durations are not reconstructed.",
        ],
    }
    write_json(output, report)
    return report
